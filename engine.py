# Calcolo pubblico SpettriNTC, derivato dalle routine di calcolo relazionepro.
# Nessuna funzionalita di generazione di relazioni, licenze o contenuto Word PRO.
import numpy as np
import re,unicodedata,math,json,threading,sys,os,urllib.request,urllib.parse,hashlib,base64,zlib
from pathlib import Path
from math import sqrt,pi,log,sin,cos,asin,radians
from pyproj import Transformer
APP_CACHE_DIR=Path.home()/".generatore_spettri_sismici_web"
GEOCODE_CACHE_FILE=APP_CACHE_DIR/"geocode_comuni_v2.json"
_GEOCODE_CACHE=None
_GEOCODE_CACHE_LOCK=threading.RLock()
def resource_path(name):
    return Path(__file__).resolve().parent/'data'/name
G = 9.81

EARTH_RADIUS_KM = 6371.0088

STATE_ORDER = ["SLO", "SLD", "SLV", "SLC"]

TR_REFERENCE = np.array([30.0, 50.0, 72.0, 101.0, 140.0, 201.0, 475.0, 975.0, 2475.0])

PVR = {"SLO": 0.81, "SLD": 0.63, "SLV": 0.10, "SLC": 0.05}

CLASS_USE = {
    "I": 0.70,
    "II": 1.00,
    "III": 1.50,
    "IV": 2.00,
}

CLASS_LABELS = {
    "I": "I - Presenza solo occasionale di persone / costruzioni agricole",
    "II": "II - Normali affollamenti e costruzioni ordinarie",
    "III": "III - Affollamenti significativi / attività rilevanti",
    "IV": "IV - Funzioni pubbliche o strategiche importanti",
}

TABLE2_ISLANDS = {
    "G1": {
        30:(0.186,2.61,0.273), 50:(0.235,2.67,0.296), 72:(0.274,2.70,0.303),
        101:(0.314,2.73,0.307), 140:(0.351,2.78,0.313), 201:(0.393,2.82,0.322),
        475:(0.500,2.88,0.340), 975:(0.603,2.98,0.372), 2475:(0.747,3.09,0.401),
    },
    "G2": {
        30:(0.239,2.61,0.245), 50:(0.303,2.61,0.272), 72:(0.347,2.61,0.298),
        101:(0.389,2.66,0.326), 140:(0.430,2.69,0.366), 201:(0.481,2.71,0.401),
        475:(0.600,2.92,0.476), 975:(0.707,3.07,0.517), 2475:(0.852,3.27,0.564),
    },
    "G3": {
        30:(0.429,2.50,0.400), 50:(0.554,2.50,0.400), 72:(0.661,2.50,0.400),
        101:(0.776,2.50,0.400), 140:(0.901,2.50,0.400), 201:(1.056,2.50,0.400),
        475:(1.500,2.50,0.400), 975:(1.967,2.50,0.400), 2475:(2.725,2.50,0.400),
    },
    "G4": {
        30:(0.350,2.70,0.400), 50:(0.558,2.70,0.400), 72:(0.807,2.70,0.400),
        101:(1.020,2.70,0.400), 140:(1.214,2.70,0.400), 201:(1.460,2.70,0.400),
        475:(2.471,2.70,0.400), 975:(3.212,2.70,0.400), 2475:(4.077,2.70,0.400),
    },
    "G5": {
        30:(0.618,2.45,0.287), 50:(0.817,2.48,0.290), 72:(0.983,2.51,0.294),
        101:(1.166,2.52,0.290), 140:(1.354,2.56,0.290), 201:(1.580,2.56,0.292),
        475:(2.200,2.58,0.306), 975:(2.823,2.65,0.316), 2475:(3.746,2.76,0.324),
    },
}

TABLE2_GROUP_LABEL = {
    "G1": "Arcipelago Toscano / Egadi / Pantelleria / Sardegna / Lampedusa-Linosa / Ponza-Palmarola-Zannone",
    "G2": "Ventotene / Santo Stefano",
    "G3": "Ustica / Tremiti",
    "G4": "Alicudi / Filicudi",
    "G5": "Panarea / Stromboli / Lipari / Vulcano / Salina",
}

ISTAT_COMUNI_XLSX_URL = (
    "https://www.istat.it/storage/codici-unita-amministrative/Elenco-comuni-italiani.xlsx"
)

ISTAT_COMUNI_XLSX_URLS = (
    ISTAT_COMUNI_XLSX_URL,
    "https://www.istat.it/wp-content/uploads/2024/09/Elenco-comuni-italiani.xlsx",
)

SARDINIA_2026_PROVINCES = {
    "112": "Città metropolitana di Sassari",
    "113": "Gallura Nord-Est Sardegna",
    "114": "Nuoro",
    "115": "Oristano",
    "116": "Ogliastra",
    "117": "Medio Campidano",
    "118": "Città metropolitana di Cagliari",
    "119": "Sulcis Iglesiente",
}

def _municipality_cache_key(row):
    code = str((row or {}).get("istat", "") or "").strip()
    code = re.sub(r"\.0$", "", code)
    if code.isdigit() and len(code) <= 6:
        code = code.zfill(6)
    if code:
        return "istat:" + code
    return "name:" + "|".join(normalize_text((row or {}).get(k, ""))
                                for k in ("regione", "provincia", "comune"))

def _load_geocode_cache():
    global _GEOCODE_CACHE
    with _GEOCODE_CACHE_LOCK:
        if _GEOCODE_CACHE is not None:
            return _GEOCODE_CACHE
        try:
            data = json.loads(GEOCODE_CACHE_FILE.read_text(encoding="utf-8"))
            _GEOCODE_CACHE = data if isinstance(data, dict) else {}
        except (OSError, ValueError, TypeError):
            _GEOCODE_CACHE = {}
        return _GEOCODE_CACHE

def _geocode_cache_get(row):
    key = _municipality_cache_key(row)
    data = _load_geocode_cache().get(key)
    try:
        lat = float(data["lat"]); lon = float(data["lon"])
        if 34.0 <= lat <= 48.5 and 5.0 <= lon <= 20.5:
            return lat, lon
    except (TypeError, ValueError, KeyError):
        pass
    return None

def _geocode_cache_put(row, lat, lon):
    try:
        lat = float(lat); lon = float(lon)
        if not (34.0 <= lat <= 48.5 and 5.0 <= lon <= 20.5):
            return
        APP_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        key = _municipality_cache_key(row)
        with _GEOCODE_CACHE_LOCK:
            cache = _load_geocode_cache()
            cache[key] = {
                "lat": round(lat, 8), "lon": round(lon, 8),
                "comune": str((row or {}).get("comune", "")),
                "provincia": str((row or {}).get("provincia", "")),
                "regione": str((row or {}).get("regione", "")),
            }
            tmp = GEOCODE_CACHE_FILE.with_suffix(".tmp")
            tmp.write_text(json.dumps(cache, ensure_ascii=False, sort_keys=True), encoding="utf-8")
            tmp.replace(GEOCODE_CACHE_FILE)
    except Exception:
        # La cache è solo un'ottimizzazione: non deve mai bloccare il calcolo.
        return

def normalize_text(value):
    value = str(value or "").strip().lower()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.replace("’", "'")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())

def haversine_km(lat1, lon1, lat2, lon2):
    """Distanza haversine; lat2/lon2 possono essere array numpy."""
    lat1r = np.radians(lat1)
    lon1r = np.radians(lon1)
    lat2r = np.radians(lat2)
    lon2r = np.radians(lon2)
    dlat = lat2r - lat1r
    dlon = lon2r - lon1r
    a = np.sin(dlat / 2.0) ** 2 + np.cos(lat1r) * np.cos(lat2r) * np.sin(dlon / 2.0) ** 2
    a = np.clip(a, 0.0, 1.0)
    return 2.0 * EARTH_RADIUS_KM * np.arcsin(np.sqrt(a))

class MunicipalityDB:
    """
    Usa esclusivamente Regione/Provincia/Comune dell'elenco ufficiale ISTAT.
    Le coordinate NON vengono lette dal catalogo dei Comuni: vengono risolte via
    Nominatim e poi memorizzate nella cache locale per rendere le selezioni successive immediate.

    Dal 2026 l'assetto territoriale sardo è stato completamente ricodificato da ISTAT:
    il parser legge TUTTI i fogli del workbook e non si arresta più al primo foglio valido.
    """
    def __init__(self):
        self.rows = []
        self.regions = []
        self._normalized_names = {}
        self.path = self._ensure_database()
        self._load()

        # Difesa contro file locali/cache vecchi o incompleti. Dal 2026 tutti i
        # 377 Comuni sardi sono stati ricodificati con prefissi 112..119: se
        # manca la Sardegna oppure troviamo solo la vecchia codifica, proviamo
        # una sola volta a recuperare l'elenco ISTAT corrente.
        sardegna_rows = [r for r in self.rows if normalize_text(r.get("regione")) == "sardegna"]
        sardegna_2026 = any(
            self._clean_istat_code(r.get("istat", ""))[:3] in SARDINIA_2026_PROVINCES
            for r in sardegna_rows
        )
        if not sardegna_rows or not sardegna_2026:
            try:
                refreshed = self._download_current_database(force=True)
                self.rows = []
                self.regions = []
                self._normalized_names = {}
                self.path = refreshed
                self._load()
            except Exception:
                # Non rendere inutilizzabile l'intera app se il PC è offline.
                pass

    @staticmethod
    def _download_current_database(force=False):
        APP_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        cached = APP_CACHE_DIR / "Elenco-comuni-italiani.xlsx"
        if not force and cached.exists() and cached.stat().st_size > 100_000:
            return cached
        last_exc = None
        for url in ISTAT_COMUNI_XLSX_URLS:
            req = urllib.request.Request(
                url,
                headers={
                    "User-Agent": "GiulivoIngegneria-SpettriNTC-PRO/3.1",
                    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
                },
            )
            try:
                with urllib.request.urlopen(req, timeout=20) as response:
                    content = response.read()
                if len(content) < 100_000 or not content.startswith(b"PK"):
                    raise RuntimeError("download ISTAT incompleto o non XLSX")
                tmp = cached.with_suffix(".tmp")
                tmp.write_bytes(content)
                tmp.replace(cached)
                return cached
            except Exception as exc:
                last_exc = exc
        raise RuntimeError(f"Download elenco ISTAT non riuscito: {last_exc}")

    def _ensure_database(self):
        # Se in una precedente esecuzione abbiamo scaricato un elenco ISTAT
        # aggiornato, preferiscilo a un eventuale file distribuito con l'app.
        cached = APP_CACHE_DIR / "Elenco-comuni-italiani.xlsx"
        if cached.exists() and cached.stat().st_size > 100_000:
            return cached
        bundled = resource_path("Elenco-comuni-italiani.xlsx")
        if bundled.exists() and bundled.stat().st_size > 100_000:
            return bundled
        try:
            return self._download_current_database(force=False)
        except Exception as exc:
            raise RuntimeError(
                "Impossibile caricare l'elenco ufficiale ISTAT dei Comuni.\n\n"
                "La connessione Internet serve al primo avvio se il file non è distribuito con l'app; "
                "il calcolo sismico resta utilizzabile inserendo direttamente le coordinate WGS84.\n\n"
                f"Dettaglio: {exc}"
            )

    @staticmethod
    def _find_header_index(headers, required_words, preferred_words=()):
        norm = [normalize_text(x) for x in headers]
        candidates = []
        for i, h in enumerate(norm):
            if all(w in h for w in required_words):
                score = sum(1 for w in preferred_words if w in h)
                # A parità di pertinenza scegli la colonna più a sinistra.
                # È importante per il file ISTAT 2026: sia il Comune sia la
                # UTS/provincia possono contenere le parole "denominazione" e
                # "italiano", ma la colonna del Comune viene prima.
                candidates.append((score, -i, i))
        if not candidates:
            return None
        return max(candidates)[2]

    @staticmethod
    def _clean_istat_code(value):
        value = str(value or "").strip()
        value = re.sub(r"\.0$", "", value)
        if value.isdigit() and len(value) <= 6:
            value = value.zfill(6)
        return value

    @staticmethod
    def _sardinia_province_from_code(code):
        code = MunicipalityDB._clean_istat_code(code)
        return SARDINIA_2026_PROVINCES.get(code[:3], "") if len(code) >= 3 else ""

    def _load(self):
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise RuntimeError(
                "Per leggere l'elenco ISTAT installare openpyxl:\n\n"
                "pip install openpyxl\n\n"
                f"Dettaglio: {exc}"
            )

        wb = load_workbook(self.path, read_only=True, data_only=True)
        loaded_sheets = 0
        try:
            for ws in wb.worksheets:
                raw_rows = ws.iter_rows(values_only=True)
                header = None
                # Alcune edizioni ISTAT hanno righe informative prima dell'intestazione.
                for _ in range(60):
                    try:
                        candidate = next(raw_rows)
                    except StopIteration:
                        break
                    norm = [normalize_text(x) for x in candidate]
                    has_region = any("denominazione regione" in x for x in norm)
                    has_comune = any(
                        ("denominazione" in x and "comune" in x) or
                        "denominazione in italiano" in x
                        for x in norm
                    )
                    if has_region and has_comune:
                        header = list(candidate)
                        break
                if header is None:
                    continue

                i_region = self._find_header_index(header, ["denominazione", "regione"])
                i_comune = self._find_header_index(header, ["denominazione", "italiano"])
                if i_comune is None:
                    i_comune = self._find_header_index(
                        header, ["denominazione"], ["italiana", "straniera", "comune"]
                    )
                i_prov = self._find_header_index(
                    header, ["denominazione", "unita", "territoriale", "sovracomunale"], ["italiano"]
                )
                if i_prov is None:
                    i_prov = self._find_header_index(header, ["denominazione", "provincia"], ["italiano"])
                i_code = self._find_header_index(header, ["codice", "comune"], ["alfanumerico"])

                # Regione e Comune sono indispensabili; la Provincia può essere
                # ricostruita per i codici sardi 2026 se il foglio la lascia vuota.
                if i_region is None or i_comune is None:
                    continue

                sheet_rows = 0
                for raw in raw_rows:
                    vals = list(raw)

                    def cell(i):
                        if i is None or i >= len(vals) or vals[i] is None:
                            return ""
                        return str(vals[i]).strip()

                    comune = cell(i_comune)
                    regione = cell(i_region)
                    istat = self._clean_istat_code(cell(i_code))
                    provincia = cell(i_prov)

                    if normalize_text(regione) == "sardegna":
                        # L'assetto 2026 ha ricodificato tutti i 377 Comuni sardi.
                        # Se la denominazione UTS non è presente/leggibile, usa il
                        # prefisso ufficiale del nuovo codice comunale.
                        provincia = provincia or self._sardinia_province_from_code(istat)

                    if comune and regione and provincia:
                        self.rows.append({
                            "comune": comune,
                            "provincia": provincia,
                            "regione": regione,
                            "istat": istat,
                        })
                        sheet_rows += 1

                if sheet_rows:
                    loaded_sheets += 1
        finally:
            wb.close()

        if not self.rows or loaded_sheets == 0:
            raise RuntimeError("Formato dell'elenco ISTAT non riconosciuto.")

        # Deduplica conservando i nomi ufficiali correnti. Importante: NON fermarsi
        # al primo foglio, perché le edizioni 2026 possono separare dati/assetti.
        unique = {}
        for r in self.rows:
            key = (normalize_text(r["regione"]), normalize_text(r["provincia"]), normalize_text(r["comune"]))
            unique[key] = r
        self.rows = sorted(unique.values(), key=lambda x: (
            normalize_text(x["regione"]), normalize_text(x["provincia"]), normalize_text(x["comune"])
        ))
        self.regions = sorted({r["regione"] for r in self.rows}, key=normalize_text)

        for i, row in enumerate(self.rows):
            self._normalized_names.setdefault(normalize_text(row["comune"]), []).append(i)

    def provinces(self, region):
        nreg = normalize_text(region)
        return sorted(
            {r["provincia"] for r in self.rows if normalize_text(r["regione"]) == nreg},
            key=normalize_text,
        )

    def municipalities(self, region, province):
        nreg = normalize_text(region)
        nprov = normalize_text(province)
        return sorted(
            [r["comune"] for r in self.rows
             if normalize_text(r["regione"]) == nreg and normalize_text(r["provincia"]) == nprov],
            key=normalize_text,
        )

    def get(self, region, province, municipality):
        nreg = normalize_text(region)
        nprov = normalize_text(province)
        nmun = normalize_text(municipality)
        for r in self.rows:
            if (normalize_text(r["regione"]) == nreg and
                    normalize_text(r["provincia"]) == nprov and
                    normalize_text(r["comune"]) == nmun):
                return r
        return None

    def find_by_name(self, municipality, region_hint=None):
        ids = self._normalized_names.get(normalize_text(municipality), [])
        if not ids:
            return None
        if region_hint:
            nreg = normalize_text(region_hint)
            for i in ids:
                rowreg = normalize_text(self.rows[i]["regione"])
                if nreg in rowreg or rowreg in nreg:
                    return self.rows[i]
        return self.rows[ids[0]]

def _near_point(lat, lon, center_lat, center_lon, radius_km):
    return float(haversine_km(lat, lon, np.array([center_lat]), np.array([center_lon]))[0]) <= radius_km

def detect_table2_island_group(lat_wgs, lon_wgs, region="", municipality=""):
    """Riconosce le isole dell'Allegato B - Tabella 2 NTC."""
    reg = normalize_text(region)
    mun = normalize_text(municipality)

    # Sardegna: riconoscimento ridondante per nome E per coordinate.
    # È importante usare anche le coordinate: quando il sito viene scelto
    # direttamente sulla mappa, Regione/Provincia/Comune vengono temporaneamente
    # svuotati prima del reverse-geocoding. Senza questo controllo la Sardegna
    # finirebbe erroneamente nel reticolo NTC, che non la copre.
    if "sardegna" in reg or "sardigna" in reg:
        return "G1"

    # Geofence WGS84 prudente dell'isola principale. Copre l'intera Sardegna
    # (circa 38.86–41.31 N, 8.13–9.84 E) con un piccolo margine numerico.
    # Il limite nord è mantenuto sotto la Corsica per evitare falsi positivi.
    try:
        _lat = float(lat_wgs)
        _lon = float(lon_wgs)
        if 38.70 <= _lat <= 41.33 and 7.95 <= _lon <= 9.95:
            return "G1"
    except (TypeError, ValueError):
        pass
    if mun in {
        "capraia isola", "campo nell elba", "capoliveri", "marciana", "marciana marina",
        "porto azzurro", "portoferraio", "rio", "isola del giglio", "favignana",
        "pantelleria", "lampedusa e linosa", "ponza"
    }:
        return "G1"
    if mun == "ventotene":
        return "G2"
    if mun in {"ustica", "isole tremiti"}:
        return "G3"
    if mun in {"leni", "malfa", "santa marina salina"}:
        return "G5"

    # Geofencing per isole che appartengono amministrativamente a comuni più grandi.
    groups = {
        "G1": [
            (42.79,10.28,35), (42.36,10.91,18), (43.04,9.84,18), (43.43,9.90,12),
            (42.33,10.31,12), (42.59,10.08,16), (42.25,11.10,12),
            (37.93,12.33,18), (37.97,12.05,16), (36.79,11.99,22),
            (35.50,12.60,22), (35.86,12.86,13), (40.91,12.97,18),
            (40.93,12.86,10), (40.97,13.05,10),
        ],
        "G2": [(40.79,13.43,15)],
        "G3": [(38.71,13.18,15), (42.12,15.50,16)],
        "G4": [(38.55,14.35,12), (38.57,14.57,13)],
        "G5": [(38.63,15.06,13), (38.79,15.21,14), (38.48,14.95,16), (38.40,14.96,14), (38.56,14.84,15)],
    }
    for group in ("G4", "G5", "G2", "G3", "G1"):
        for clat, clon, rad in groups[group]:
            if _near_point(lat_wgs, lon_wgs, clat, clon, rad):
                return group
    return None

class SeismicHazardDB:
    """
    Reticolo ufficiale NTC (10.751 nodi, coordinate ED50) + Allegato B Tabella 2 isole.
    Nel file ufficiale ag è espresso in g/10; l'output dell'app è ag/g.
    """

    def __init__(self, csv_path, dense_ag_path=None):
        csv_path = Path(csv_path)
        if not csv_path.exists():
            raise FileNotFoundError(
                f"Database sismico non trovato:\n{csv_path}\n\n"
                "Estrai l'intero ZIP mantenendo la cartella data/ accanto a relazionepro.py. "
                "Non copiare soltanto il file Python."
            )

        self.data = np.genfromtxt(csv_path, delimiter=",", names=True, dtype=float, encoding="utf-8")
        self.lat = np.asarray(self.data["LAT"], dtype=float)
        self.lon = np.asarray(self.data["LON"], dtype=float)
        self.ids = np.asarray(self.data["ID"], dtype=int)

        # Il reticolo dell'Allegato B deriva da una griglia logica di 222 colonne.
        # Conserviamo la topologia ID -> (riga, colonna) per individuare la MAGLIA
        # elementare reale contenente il sito, invece di scegliere semplicemente
        # quattro punti vicini. Questo evita errori soprattutto presso coste/confini.
        self.grid_rows = (self.ids - 1) // 222
        self.grid_cols = (self.ids - 1) % 222
        self.grid_index = {
            (int(r), int(c)): int(i)
            for i, (r, c) in enumerate(zip(self.grid_rows, self.grid_cols))
        }

        self.dense_lon = self.dense_lat = self.dense_ag = None
        dense_ag_path = Path(dense_ag_path) if dense_ag_path else None
        if dense_ag_path and dense_ag_path.exists():
            try:
                dense = np.load(dense_ag_path)
                self.dense_lon = np.asarray(dense["lon"], dtype=float)
                self.dense_lat = np.asarray(dense["lat"], dtype=float)
                self.dense_ag = np.asarray(dense["ag"], dtype=float)
            except Exception:
                self.dense_lon = self.dense_lat = self.dense_ag = None

    @staticmethod
    def _bounded_tr(tr):
        return float(min(2475.0, max(30.0, float(tr))))

    def _time_bounds(self, tr):
        tr = self._bounded_tr(tr)
        if tr <= TR_REFERENCE[0]:
            return 30, 30, 0.0
        if tr >= TR_REFERENCE[-1]:
            return 2475, 2475, 0.0
        exact = np.where(np.isclose(TR_REFERENCE, tr, atol=1e-12))[0]
        if exact.size:
            t = int(TR_REFERENCE[exact[0]])
            return t, t, 0.0
        upper_i = int(np.searchsorted(TR_REFERENCE, tr, side="right"))
        lower_i = upper_i - 1
        t1, t2 = int(TR_REFERENCE[lower_i]), int(TR_REFERENCE[upper_i])
        alpha = log(tr / t1) / log(t2 / t1)
        return t1, t2, alpha

    def _node_values(self, indices, tr, parameter):
        t1, t2, alpha = self._time_bounds(tr)
        v1 = np.asarray(self.data[f"T{t1}{parameter}"][indices], dtype=float)
        if t1 == t2:
            return v1
        v2 = np.asarray(self.data[f"T{t2}{parameter}"][indices], dtype=float)
        return np.exp(np.log(v1) + np.log(v2 / v1) * alpha)

    def _table2_values(self, tr, group):
        if group not in TABLE2_ISLANDS:
            raise ValueError("Gruppo isola Tabella 2 non valido.")
        t1, t2, alpha = self._time_bounds(tr)
        a1, f1, c1 = TABLE2_ISLANDS[group][t1]
        if t1 == t2:
            return a1, f1, c1
        a2, f2, c2 = TABLE2_ISLANDS[group][t2]
        interp = lambda v1, v2: float(np.exp(np.log(v1) + np.log(v2 / v1) * alpha))
        return interp(a1,a2), interp(f1,f2), interp(c1,c2)

    @staticmethod
    def _point_in_polygon(x, y, polygon):
        """Point-in-polygon robusto per la piccola maglia quadrilatera del reticolo."""
        inside = False
        n = len(polygon)
        j = n - 1
        for i in range(n):
            xi, yi = polygon[i]
            xj, yj = polygon[j]
            # Punto sul segmento: consideralo interno.
            dx, dy = xj - xi, yj - yi
            cross = (x - xi) * dy - (y - yi) * dx
            if abs(cross) <= 1e-10:
                dot = (x - xi) * (xj - xi) + (y - yi) * (yj - yi)
                seg2 = dx * dx + dy * dy
                if -1e-12 <= dot <= seg2 + 1e-12:
                    return True
            intersects = ((yi > y) != (yj > y)) and (
                x < (xj - xi) * (y - yi) / ((yj - yi) if abs(yj - yi) > 1e-15 else 1e-15) + xi
            )
            if intersects:
                inside = not inside
            j = i
        return inside

    def _cell_vertices(self, lat_ed50, lon_ed50):
        """
        Restituisce i QUATTRO vertici della maglia elementare reale contenente il sito.

        Gli ID del reticolo NTC seguono la griglia logica originaria: avanzare di 1
        significa passare al nodo adiacente lungo una direzione della griglia, avanzare
        di 222 significa passare alla riga adiacente. Cerchiamo quindi le celle complete
        attorno ai nodi più vicini e verifichiamo geometricamente quale contiene il punto.
        """
        distances = haversine_km(lat_ed50, lon_ed50, self.lat, self.lon)
        n_candidates = min(24, len(distances))
        nearest = np.argpartition(distances, n_candidates - 1)[:n_candidates]
        nearest = nearest[np.argsort(distances[nearest])]

        candidates = []
        visited_cells = set()
        for idx in nearest:
            r = int(self.grid_rows[idx])
            c = int(self.grid_cols[idx])
            # Il nodo può appartenere a quattro celle; proviamole tutte.
            for r0 in (r - 1, r):
                for c0 in (c - 1, c):
                    key = (r0, c0)
                    if key in visited_cells:
                        continue
                    visited_cells.add(key)
                    keys = ((r0, c0), (r0, c0 + 1), (r0 + 1, c0 + 1), (r0 + 1, c0))
                    try:
                        verts = np.array([self.grid_index[k] for k in keys], dtype=int)
                    except KeyError:
                        continue
                    poly = [(float(self.lon[i]), float(self.lat[i])) for i in verts]
                    if self._point_in_polygon(float(lon_ed50), float(lat_ed50), poly):
                        candidates.append((float(np.sum(distances[verts])), verts))

        if candidates:
            candidates.sort(key=lambda item: item[0])
            return candidates[0][1], distances

        # Vicinissimo a un nodo esatto la cella può essere ambigua sul bordo: il valore
        # nel nodo è comunque definito e site_parameters lo gestirà con peso unitario.
        if float(distances[nearest[0]]) < 1e-6:
            return np.array([int(nearest[0])], dtype=int), distances

        # Non inventare una maglia usando quattro nodi qualsiasi. Se siamo molto vicini
        # al reticolo ma il punto è su una costa/confine, usiamo il criterio dei quattro
        # quadranti come fallback esplicito; altrimenti segnaliamo sito non coperto.
        if float(distances[nearest[0]]) <= 12.0:
            cand = nearest
            quadrants = [
                cand[(self.lat[cand] >= lat_ed50) & (self.lon[cand] <= lon_ed50)],
                cand[(self.lat[cand] >= lat_ed50) & (self.lon[cand] >= lon_ed50)],
                cand[(self.lat[cand] <= lat_ed50) & (self.lon[cand] <= lon_ed50)],
                cand[(self.lat[cand] <= lat_ed50) & (self.lon[cand] >= lon_ed50)],
            ]
            chosen = []
            for q in quadrants:
                if q.size:
                    chosen.append(int(q[np.argmin(distances[q])]))
            if len(set(chosen)) == 4:
                return np.array(chosen, dtype=int), distances

        raise ValueError(
            "Non è stato possibile individuare in modo univoco la maglia elementare NTC "
            "contenente il punto. Verificare le coordinate del sito o l'eventuale appartenenza "
            "alle isole dell'Allegato B - Tabella 2."
        )

    def site_parameters(self, lat_ed50, lon_ed50, tr, table2_group=None):
        if table2_group:
            ag_table, f0, tc = self._table2_values(tr, table2_group)
            return {
                "Tr": self._bounded_tr(tr), "ag_g": ag_table / 10.0, "F0": f0, "Tc_star": tc,
                "node_ids": [], "node_distances_km": [], "source": f"Tabella 2 - {TABLE2_GROUP_LABEL[table2_group]}",
            }

        vertices, distances = self._cell_vertices(lat_ed50, lon_ed50)
        d = distances[vertices]
        # Se il sito non è in una zona coperta dal reticolo e non è stato riconosciuto
        # come isola Tabella 2, non restituiamo valori falsi da nodi lontani.
        if float(np.min(d)) > 35.0:
            raise ValueError(
                "Il punto è esterno alla copertura del reticolo NTC e non è stato riconosciuto "
                "come isola dell'Allegato B - Tabella 2. Verificare le coordinate del sito."
            )
        if np.min(d) < 1e-6:
            i = int(vertices[np.argmin(d)])
            vertices = np.array([i], dtype=int)
            weights = np.array([1.0])
        else:
            weights = 1.0 / d
            weights = weights / weights.sum()

        ag_table = float(np.sum(self._node_values(vertices, tr, "ag") * weights))
        f0 = float(np.sum(self._node_values(vertices, tr, "F0") * weights))
        tc = float(np.sum(self._node_values(vertices, tr, "Tc") * weights))
        return {
            "Tr": self._bounded_tr(tr), "ag_g": ag_table / 10.0, "F0": f0, "Tc_star": tc,
            "node_ids": [int(self.ids[i]) for i in vertices],
            "node_distances_km": [float(x) for x in distances[vertices]],
            "source": "Reticolo NTC - Allegato B Tabella 1",
        }

    def hazard_points(self, lat_ed50, lon_ed50, tr, radius_km=90.0, table2_group=None):
        tr = self._bounded_tr(tr)
        if self.dense_lon is not None and self.dense_lat is not None and self.dense_ag is not None and abs(tr - 475.0) <= 0.5:
            d = haversine_km(lat_ed50, lon_ed50, self.dense_lat, self.dense_lon)
            mask = d <= radius_km
            if int(np.sum(mask)) >= 4:
                return self.dense_lon[mask], self.dense_lat[mask], self.dense_ag[mask], "INGV - griglia ag 0.02° (Tr=475 anni)"

        if table2_group:
            ag_table, _, _ = self._table2_values(tr, table2_group)
            ag = ag_table / 10.0
            # Campo uniforme locale solo per visualizzare la costanza prescritta dalla Tabella 2.
            dlat = max(0.05, radius_km / 111.0)
            dlon = max(0.05, radius_km / max(40.0, 111.0 * np.cos(np.radians(lat_ed50))))
            xs = np.linspace(lon_ed50-dlon, lon_ed50+dlon, 7)
            ys = np.linspace(lat_ed50-dlat, lat_ed50+dlat, 7)
            X, Y = np.meshgrid(xs, ys)
            Z = np.full(X.size, ag, dtype=float)
            return X.ravel(), Y.ravel(), Z, f"NTC Allegato B - Tabella 2 ({TABLE2_GROUP_LABEL[table2_group]})"

        d = haversine_km(lat_ed50, lon_ed50, self.lat, self.lon)
        idx = np.where(d <= radius_km)[0]
        ag_g = self._node_values(idx, tr, "ag") / 10.0
        return self.lon[idx], self.lat[idx], ag_g, "Reticolo NTC - ag interpolata nel tempo"

def get_Ss_Cc(ag_g, F0, cat_s, Tc_star):
    """Restituisce Ss e Cc."""
    cat_s = cat_s.upper().strip()
    agFo_g = F0 * ag_g

    if cat_s == "A":
        Ss = 1.0
        Cc = 1.0
    elif cat_s == "B":
        Ss = min(1.20, max(1.00, 1.40 - 0.40 * agFo_g))
        Cc = 1.10 * Tc_star ** (-0.20)
    elif cat_s == "C":
        Ss = min(1.50, max(1.00, 1.70 - 0.60 * agFo_g))
        Cc = 1.05 * Tc_star ** (-0.33)
    elif cat_s == "D":
        Ss = min(1.80, max(0.90, 2.40 - 1.50 * agFo_g))
        Cc = 1.25 * Tc_star ** (-0.50)
    elif cat_s == "E":
        Ss = min(1.60, max(1.00, 2.00 - 1.10 * agFo_g))
        Cc = 1.15 * Tc_star ** (-0.40)
    else:
        raise ValueError("Categoria di sottosuolo non valida. Usare A, B, C, D o E.")

    return Ss, Cc

def get_St(topo):
    vals = {"T1": 1.0, "T2": 1.2, "T3": 1.2, "T4": 1.4}
    topo = topo.upper().strip()
    if topo not in vals:
        raise ValueError("Categoria topografica non valida.")
    return vals[topo]

def calculate_spectrum(ag_g, F0, Tc_star, soil="C", topo="T1",
                       xi=5.0, q=1.0, t_max=5.0, n_points=1600, state=None):
    """
    Calcola spettro elastico, spettro di progetto e spettri di spostamento.
    Accelerazioni restituite in g; spostamenti in mm.
    """

    if ag_g <= 0:
        raise ValueError("ag/g deve essere > 0.")
    if F0 <= 0:
        raise ValueError("F0 deve essere > 0.")
    if Tc_star <= 0:
        raise ValueError("Tc* deve essere > 0.")
    if xi <= 0:
        raise ValueError("Lo smorzamento xi deve essere > 0.")
    if q < 1.0:
        raise ValueError("q deve essere >= 1.0.")
    if t_max <= 0:
        raise ValueError("Tmax deve essere > 0.")

    Ss, Cc = get_Ss_Cc(ag_g, F0, soil, Tc_star)
    St = get_St(topo)
    S = Ss * St

    eta = max(0.55, sqrt(10.0 / (5.0 + xi)))
    beta = 1.0 / q

    TC = Cc * Tc_star
    TB = TC / 3.0
    TD = 4.0 * ag_g + 1.6

    T = np.linspace(0.0, t_max, n_points)
    Se_g = np.zeros_like(T)
    Sd_g = np.zeros_like(T)
    SDe_mm = np.zeros_like(T)
    SDd_mm = np.zeros_like(T)

    for i, t in enumerate(T):
        if t <= TB:
            Se = (
                G * ag_g * S * eta * F0 *
                ((t / TB) + (1.0 / (eta * F0)) * (1.0 - t / TB))
            )
        elif t <= TC:
            Se = G * ag_g * S * eta * F0
        elif t <= TD:
            Se = G * ag_g * S * eta * F0 * (TC / t)
        else:
            Se = G * ag_g * S * eta * F0 * (TC * TD) / (t ** 2)

        # NTC 2018 §3.2.3.4: allo SLO lo spettro di progetto coincide
        # con il corrispondente spettro elastico. Per SLD/SLV/SLC (§3.2.3.5)
        # la riduzione delle ordinate nelle analisi lineari è introdotta tramite q.
        if str(state or "").upper() == "SLO":
            Sd_acc = Se
        else:
            if t <= TB:
                Sd_acc = (
                    G * ag_g * S * beta * F0 *
                    ((t / TB) + (1.0 / (beta * F0)) * (1.0 - t / TB))
                )
            elif t <= TC:
                Sd_acc = G * ag_g * S * beta * F0
            elif t <= TD:
                Sd_acc = G * ag_g * S * beta * F0 * (TC / t)
            else:
                Sd_acc = G * ag_g * S * beta * F0 * (TC * TD) / (t ** 2)

            Sd_acc = max(Sd_acc, 0.2 * ag_g * G)

        Se_g[i] = Se / G
        Sd_g[i] = Sd_acc / G
        # Se e Sd_acc sono accelerazioni [m/s²]; moltiplicando per (T/2π)² si ottengono metri,
        # quindi ×1000 restituisce gli spostamenti spettrali in millimetri.
        SDe_mm[i] = Se * (t / (2.0 * pi)) ** 2 * 1000.0
        SDd_mm[i] = Sd_acc * (t / (2.0 * pi)) ** 2 * 1000.0

    return {
        "T": T,
        "Se_g": Se_g,
        "Sd_g": Sd_g,
        "Sa_g": Sd_g,  # alias storico interno, mantenuto per compatibilità
        "SDe_mm": SDe_mm,
        "SDd_mm": SDd_mm,
        "Sd_mm": SDd_mm,  # alias storico interno, mantenuto per compatibilità
        "Ss": Ss,
        "Cc": Cc,
        "St": St,
        "S": S,
        "eta": eta,
        "beta": beta,
        "TB": TB,
        "TC": TC,
        "TD": TD,
        "Se0_g": ag_g * S,
        "plateau_el_g": ag_g * S * eta * F0,
        "plateau_pr_g": (ag_g * S * eta * F0 if str(state or "").upper() == "SLO"
                           else max(ag_g * S * beta * F0, 0.2 * ag_g)),
    }
